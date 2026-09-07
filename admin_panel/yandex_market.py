"""
Экспорт выбранных книг в Excel-шаблон Яндекс Маркета
(shablon/файл с товарами.xlsx, лист «Список товаров»).

Ozon-код (views.py, admin.py) не затрагивается — вся логика Яндекса здесь.
"""
import logging
import os
import tempfile
import unicodedata

from django.conf import settings
from django.http import HttpResponse

from openpyxl import load_workbook

from .models import Book

logger = logging.getLogger(__name__)

# Путь к шаблону Яндекс Маркета
YANDEX_TEMPLATE_PATH = os.path.join(settings.BASE_DIR, 'shablon', 'файл с товарами.xlsx')

# Код ТН ВЭД для книг
YANDEX_TNVED_CODE = '4901990000'

# Максимальная длина описания в шаблоне Маркета
YANDEX_DESCRIPTION_MAX_LEN = 6000

# Маппинг book_type → «Категория на Маркете *» (см. лист Enums шаблона)
BOOK_TYPE_YANDEX_CATEGORY_MAPPING = {
    'printed book': 'Художественная литература',
    'second': 'Нехудожественная литература',
    'bookinist': 'Нехудожественная литература',
    'print_on_demand': 'Художественная литература',
}


def _yandex_category(book: Book) -> str:
    return BOOK_TYPE_YANDEX_CATEGORY_MAPPING.get(book.book_type, 'Художественная литература')


def _yandex_age_label(book: Book) -> str:
    age = book.get_age_restrictions_display() or ''
    digits = ''.join(ch for ch in age if ch.isdigit())
    if not digits:
        return ''
    n = int(digits)
    if n % 10 == 1 and n % 100 != 11:
        return f'{n} год'
    if n % 10 in (2, 3, 4) and n % 100 not in (12, 13, 14):
        return f'{n} года'
    return f'{n} лет'


def _slugify_name(name: str) -> str:
    """ASCII-safe slug для имени файла (тот же подход, что в Ozon-экспорте)."""
    normalized = unicodedata.normalize('NFKD', name)
    ascii_bytes = normalized.encode('ascii', 'ignore').decode('ascii')
    slug = ascii_bytes.lower().replace(' ', '_').replace('/', '_')
    slug = ''.join(ch for ch in slug if ch.isalnum() or ch == '_')
    return slug[:50] or 'template'

def _yandex_headers(ws):
    """Заголовки из строки 2 листа «Список товаров» (headerAddress=A2)."""
    headers = {}
    for col_num in range(1, ws.max_column + 1):
        cell_value = ws.cell(row=2, column=col_num).value
        if cell_value:
            header_clean = str(cell_value).replace('\n', ' ').strip().lower()
            headers[header_clean] = col_num
    return headers


def _yandex_field_mapping(media_base_url):
    """Маппинг: нормализованное название колонки шаблона → функция(book)."""
    return {
        'ваш sku': lambda book: book.sku or '',
        'название товара': lambda book: book.title or '',
        'ссылка на изображение': lambda book: (
            f'{media_base_url}{book.photos[0]}' if book.photos else ''
        ),
        'описание товара': lambda book: (book.description or '')[:YANDEX_DESCRIPTION_MAX_LEN],
        'категория на маркете': _yandex_category,
        'бренд': lambda book: book.publisher or 'Нет бренда',
        'штрихкод': lambda book: book.isbn_digits or book.isbn or '',
        'страна производства': lambda book: 'Россия',
        'артикул производителя': lambda book: book.sku or '',
        'вес, кг': lambda book: round(float(book.weight) / 1000, 3) if book.weight else '',
        'длина, см': lambda book: round(float(book.length) / 10, 1) if book.length else '',
        'ширина, см': lambda book: round(float(book.width) / 10, 1) if book.width else '',
        'высота, см': lambda book: round(float(book.height) / 10, 1) if book.height else '',
        'товар доставляется в нескольких коробках': lambda book: 'Нет',
        'цена': lambda book: float(book.price) if book.price else '',
        'зачёркнутая цена': lambda book: float(book.old_price) if book.old_price else '',
        'код тн вэд': lambda book: YANDEX_TNVED_CODE,
        'с какого возраста пользоваться': _yandex_age_label,
        'товар для взрослых': lambda book: 'Да' if book.is_adult == 'yes' else '',
    }


def _fill_yandex_sheet(ws, books, headers, field_mapping):
    """Заполняет книги начиная со строки 5 (строки 1–4 шаблона не трогаем)."""
    current_row = 5
    for book in books:
        for header_name, col_num in headers.items():
            base_name = header_name.rstrip('*').strip()
            mapper = (
                field_mapping.get(header_name)
                or field_mapping.get(base_name)
                or next(
                    (v for k, v in field_mapping.items() if k.rstrip('*').strip() == base_name),
                    None,
                )
            )
            if mapper:
                try:
                    value = mapper(book)
                    if value not in (None, ''):
                        ws.cell(row=current_row, column=col_num).value = value
                except Exception as e:
                    logger.warning(
                        f"Яндекс: ошибка заполнения '{header_name}' "
                        f"для книги {book.id} ({book.sku}): {e}"
                    )
        current_row += 1


def export_books_to_yandex_template(request):
    """Экспорт выбранных книг в шаблон Яндекс Маркета.

    QuerySet берётся из request.yandex_export_queryset (ставится экшеном админки),
    иначе — все книги.
    """
    books = list(getattr(request, 'yandex_export_queryset', None) or Book.objects.all())
    if not books:
        return HttpResponse(
            "Нет книг для экспорта в шаблон Яндекс Маркета.",
            content_type='text/plain; charset=utf-8',
            status=400,
        )

    # Приоритет: активный шаблон из админки (YandexTemplate) → файл по умолчанию
    from .models import YandexTemplate

    template_path = YandexTemplate.get_active_path()
    template_source = 'админка (YandexTemplate)'
    if not template_path:
        template_path = YANDEX_TEMPLATE_PATH
        template_source = 'shablon/файл с товарами.xlsx (по умолчанию)'

    if not os.path.exists(template_path):
        return HttpResponse(
            f"Шаблон Яндекс Маркета не найден: {template_path}",
            content_type='text/plain; charset=utf-8',
            status=500,
        )
    logger.info("Яндекс: используется шаблон из %s: %s", template_source, template_path)

    media_base_url = getattr(settings, 'MEDIA_BASE_URL', None) or settings.MEDIA_URL
    if not media_base_url.endswith('/'):
        media_base_url += '/'
    # Относительный MEDIA_URL ('/media/') → абсолютный URL по ходу запроса,
    # т.к. Маркет требует прямые ссылки на изображения.
    if media_base_url.startswith('/'):
        try:
            media_base_url = request.build_absolute_uri(media_base_url)
        except Exception:
            pass

    wb = load_workbook(template_path)
    if 'Список товаров' not in wb.sheetnames:
        return HttpResponse(
            "В шаблоне Яндекс Маркета нет листа 'Список товаров'.",
            content_type='text/plain; charset=utf-8',
            status=500,
        )
    ws = wb['Список товаров']
    headers = _yandex_headers(ws)
    field_mapping = _yandex_field_mapping(media_base_url)
    _fill_yandex_sheet(ws, books, headers, field_mapping)

    tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
    wb.save(tmp.name)
    tmp.close()

    with open(tmp.name, 'rb') as f:
        file_data = f.read()
    os.unlink(tmp.name)

    filename = f"yandex_market_{len(books)}_books.xlsx"
    response = HttpResponse(
        file_data,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response

